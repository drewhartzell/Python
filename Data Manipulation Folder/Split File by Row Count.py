import os
from openpyxl import load_workbook, Workbook
from copy import copy

# Define the path to the Downloads folder and the Excel file
downloads_folder = os.path.expanduser("~/Downloads")
split_folder = os.path.join(downloads_folder, "File Name")
os.makedirs(split_folder, exist_ok=True)

file_path = os.path.join(downloads_folder, "File Name.xlsx")
sheet_name = "Sheet Name"

# Load the workbook and the specified data sheet with data_only to avoid formulas
wb = load_workbook(file_path, data_only=True)
ws = wb[sheet_name]

# Extract the first 4 rows as header
header_rows = []
for row in ws.iter_rows(min_row=1, max_row=4):
    header_rows.append([copy(cell) for cell in row])

# Start processing from row 7
start_data_row = 7
rows_to_capture = 28
rows_to_skip = 2
file_counter = 1
current_row = start_data_row

# Determine the maximum number of columns
max_col = ws.max_column

# Loop through the sheet and save chunks
while current_row <= ws.max_row:
    # Create a new workbook and sheet
    new_wb = Workbook()
    new_ws = new_wb.active

    # Write header rows with formatting
    for r_idx, row in enumerate(header_rows, start=1):
        for c_idx, cell in enumerate(row, start=1):
            new_cell = new_ws.cell(row=r_idx, column=c_idx, value=cell.value)
            new_cell.font = copy(cell.font)
            new_cell.border = copy(cell.border)
            new_cell.fill = copy(cell.fill)
            new_cell.number_format = copy(cell.number_format)
            new_cell.alignment = copy(cell.alignment)

    # Write data rows with formatting and only values
    for i in range(rows_to_capture):
        source_row_idx = current_row + i
        if source_row_idx > ws.max_row:
            break
        for col_idx in range(1, max_col + 1):
            source_cell = ws.cell(row=source_row_idx, column=col_idx)
            target_cell = new_ws.cell(row=4 + i + 1, column=col_idx, value=source_cell.value)
            target_cell.font = copy(source_cell.font)
            target_cell.border = copy(source_cell.border)
            target_cell.fill = copy(source_cell.fill)
            target_cell.number_format = copy(source_cell.number_format)
            target_cell.alignment = copy(source_cell.alignment)

    # Save the new workbook with a temporary name
    temp_filename = os.path.join(split_folder, f'Temp File {file_counter}.xlsx')
    new_wb.save(temp_filename)

    # What the file will be named as - particular cell in the 'chunk' that th
    temp_wb = load_workbook(temp_filename, data_only=True)
    temp_ws = temp_wb.active
    project_name = temp_ws['C5'].value
    project_name_safe = str(project_name).replace("/", "-").replace("\\", "-") if project_name else f"Project_{file_counter}"

    # Final filename
    final_filename = os.path.join(split_folder, f'{project_name_safe}.xlsx')
    os.rename(temp_filename, final_filename)
    print(f"Saved {final_filename} with data rows {current_row} to {current_row + rows_to_capture - 1}")

    # Update counters
    current_row += rows_to_capture + rows_to_skip
    file_counter += 1
